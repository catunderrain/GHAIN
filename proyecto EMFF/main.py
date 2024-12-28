import docx
from pprint import pprint
import os
import docx.document
from openpyxl import load_workbook
from PyPDF2 import PdfReader
import sys
from datetime import datetime

SPACENUM = 30


def extract_docx_metadata(docxfile):
    doc = docx.Document(docxfile)
    core_properties = doc.core_properties
    metadata = {}
    for prop in dir(core_properties):
        if prop.startswith('___') or prop.startswith('__') or prop.startswith('_'):
            continue
        value = getattr(core_properties, prop)
        if callable(value): continue
        if prop == 'created' or prop == 'modified' or prop == 'last_printed':
            if value:
                value = value.strftime('%d-%m-%Y %H:%M:%S')
            else:
                value = None
        metadata[prop] = value
    try:
        custom_properties = core_properties.custom_properties
        if custom_properties:
            metadata['custom_properties'] = {}
            for prop in custom_properties:
                metadata['custom_properties'][prop.name] = prop.value
    except AttributeError:
        pass
    return metadata


def extract_xlsx_metadata(xlsxfile):
    wb = load_workbook(xlsxfile, read_only=None)
    props = wb.properties
    metadata = {}
    for prop in dir(props):
        if prop.startswith('___') or prop.startswith('__') or prop.startswith('_'):
            continue
        value = getattr(props, prop)
        if callable(value): continue
        if prop == 'created' or prop == 'modified' or prop == 'last_printed':
                if value:
                    value = value.strftime('%d-%m-%Y %H:%M:%S')
                else:
                    value = None
        metadata[prop] = value
    return metadata


def extract_pdf_metadata(pdffile):
    reader = PdfReader(pdffile)
    metadata = reader.metadata
    metadatanew = {}
    for k, v in metadata.items():
        metadatanew[k] = v
    if '/CreationDate' in metadata:
        creation_date = metadata['/CreationDate']
        date_str = creation_date[2:].split('+')[0]
        date_obj = datetime.strptime(date_str, '%Y%m%d%H%M%S')
        metadatanew['/CreationDate'] = date_obj.strftime('%d-%m-%Y %H:%M:%S')
    return metadatanew


def main():

    try:
        sys.argv[1]
        file_path = sys.argv[1]
    except:
        print('No file path included')
        exit()

    file_extension = file_path.split('.')[-1]

    if file_extension == 'docx':
        metadata = extract_docx_metadata(file_path)
        metadata = dict(sorted(metadata.items(), key=lambda item: str(item[1])))
        for key, val in metadata.items():
            if not val == '':
                if val is not None:
                    print(f'{key}{(SPACENUM-len(key))*' '}: {val}')
    elif file_extension == 'pdf':
        metadata = extract_pdf_metadata(file_path)
        metadata = dict(sorted(metadata.items(), key=lambda item: str(item[1])))
        for key, val in metadata.items():
            if not val == '':
                if val is not None:
                    print(f'{key}{(SPACENUM-len(key))*' '}: {val}')
    elif file_extension == 'xlsx':
        metadata = extract_xlsx_metadata(file_path)
        metadata = dict(sorted(metadata.items(), key=lambda item: str(item[1])))
        for key, val in metadata.items():
            if not val == '':
                if val is not None:
                    print(f'{key}{(SPACENUM-len(key))*' '}: {val}')
    else:
        print(f'File extension \'{file_extension}\' does not supported')
        

if __name__ == '__main__':
    main()